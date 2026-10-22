#!/usr/bin/env python3
"""Find the current Biedronka grocery leaflet and create a review-first workbook."""
from __future__ import annotations

import hashlib
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
INDEX_URL = "https://iulotka.pl/biedronka"
UA = "Mozilla/5.0 (compatible; GazetkiDoExcela/1.0; +https://github.com/)"
SESSION = requests.Session()
SESSION.headers.update({"User-Agent": UA, "Accept-Language": "pl-PL,pl;q=0.9,en;q=0.5"})


def log(msg: str) -> None:
    print(msg, flush=True)


def fetch(url: str) -> tuple[str, str]:
    response = SESSION.get(url, timeout=45)
    response.raise_for_status()
    return response.text, response.url


def find_leaflet() -> tuple[str, str, str]:
    page, final_url = fetch(INDEX_URL)
    soup = BeautifulSoup(page, "html.parser")
    found: list[tuple[int, str, str]] = []
    for order, anchor in enumerate(soup.find_all("a", href=True)):
        href = urljoin(final_url, anchor.get("href", ""))
        title = " ".join(anchor.get_text(" ", strip=True).split())
        descriptor = (title + " " + href).lower()
        if "biedronka" in descriptor and ("gazetka-spozywcza" in descriptor or "gazetka spożywcza" in title.lower()):
            found.append((order, href, title or "Biedronka — gazetka spożywcza"))
    if not found:
        raise RuntimeError("Nie znaleziono gazetki spożywczej Biedronki na stronie iUlotka.")
    # The first matching link is the newest article in iUlotka's current-feed page.
    _, detail_url, title = min(found, key=lambda item: item[0])
    return detail_url, title, final_url


def find_pdf(detail_url: str) -> str:
    detail, final_url = fetch(detail_url)
    soup = BeautifulSoup(detail, "html.parser")
    candidates = []
    for anchor in soup.find_all("a", href=True):
        label = " ".join(anchor.get_text(" ", strip=True).split()).lower()
        url = urljoin(final_url, anchor["href"])
        if "pdf" in label or urlparse(url).path.lower().endswith(".pdf"):
            candidates.append(url)
    if not candidates:
        raise RuntimeError("Nie znaleziono linku „Wersja PDF” w artykule iUlotka.")
    for url in candidates:
        try:
            response = SESSION.get(url, timeout=60, stream=True)
            response.raise_for_status()
            content_type = response.headers.get("Content-Type", "").lower()
            first = next(response.iter_content(8), b"")
            response.close()
            if first.startswith(b"%PDF") or "pdf" in content_type:
                return url
        except requests.RequestException:
            continue
    raise RuntimeError("Link PDF z iUlotka nie zwrócił pliku PDF do pobrania.")


def download_pdf(url: str) -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "biedronka-gazetka.pdf"
    response = SESSION.get(url, timeout=120, stream=True)
    response.raise_for_status()
    with path.open("wb") as stream:
        for chunk in response.iter_content(1024 * 1024):
            if chunk:
                stream.write(chunk)
    if path.stat().st_size < 20_000 or path.read_bytes()[:4] != b"%PDF":
        path.unlink(missing_ok=True)
        raise RuntimeError("Pobrany materiał nie wygląda na poprawny plik PDF.")
    return path


PRICE_RE = re.compile(r"(?<!\d)(\d{1,3})[,.](\d{2})(?!\d)")
SPACED_PRICE_RE = re.compile(r"(?<!\d)(\d{1,2})\s+(\d{2})(?!\d)")
DATE_RE = re.compile(r"\b\d{1,2}[.]\d{1,2}\s*(?:-|–|do)\s*\d{1,2}[.]\d{1,2}\b", re.I)


def ocr_lines(image, lang: str, config: str) -> list[dict]:
    """Return spatially grouped OCR lines with Tesseract word confidence."""
    import pytesseract

    data = pytesseract.image_to_data(image, lang=lang, config=config, output_type=pytesseract.Output.DICT)
    grouped: dict[tuple[int, int, int], list[int]] = {}
    for index, word in enumerate(data["text"]):
        if not word.strip():
            continue
        key = (data["block_num"][index], data["par_num"][index], data["line_num"][index])
        grouped.setdefault(key, []).append(index)
    lines = []
    for indexes in grouped.values():
        words = [data["text"][i].strip() for i in indexes]
        confs = [float(data["conf"][i]) for i in indexes if str(data["conf"][i]).replace("-", "", 1).isdigit() and float(data["conf"][i]) >= 0]
        x0 = min(data["left"][i] for i in indexes)
        y0 = min(data["top"][i] for i in indexes)
        x1 = max(data["left"][i] + data["width"][i] for i in indexes)
        y1 = max(data["top"][i] + data["height"][i] for i in indexes)
        lines.append({"text": " ".join(words), "confidence": sum(confs) / len(confs) if confs else 0.0,
                      "x0": x0, "x1": x1, "y0": y0, "y1": y1, "cx": (x0 + x1) / 2, "cy": (y0 + y1) / 2})
    return sorted(lines, key=lambda line: (line["y0"], line["x0"]))


def line_candidates(lines, page_no: int, image_width: int) -> list[dict]:
    candidates = []
    for line in lines:
        raw = re.sub(r"\s+", " ", line["text"]).strip()
        cleaned = DATE_RE.sub(" ", raw)
        if re.search(r"(?:/\s*(?:kg|l|100\s*g|100\s*ml)|za\s*kg|za\s*litr|cena\s*za\s*(?:kg|l))", cleaned, re.I):
            continue
        matches = list(PRICE_RE.finditer(cleaned))
        spaced_price = False
        if not matches and re.search(r"(?:zł|zl|pln|z1)", cleaned, re.I):
            # OCR sometimes splits the cents into a separate token ("14 99").
            # Only accept this ambiguous form with a currency marker and mark it low.
            matches = list(SPACED_PRICE_RE.finditer(cleaned))
            spaced_price = bool(matches)
        prices = []
        for match in matches:
            amount = float(f"{match.group(1)}.{match.group(2)}")
            if 0.01 <= amount <= 999.99:
                prices.append(amount)
        if not prices:
            continue
        nearby = []
        for other in lines:
            if other is line or abs(other["cy"] - line["cy"]) > 190:
                continue
            overlap = max(0, min(line["x1"], other["x1"]) - max(line["x0"], other["x0"]))
            narrow = min(max(1, line["x1"] - line["x0"]), max(1, other["x1"] - other["x0"]))
            center_distance = abs(other["cx"] - line["cx"])
            if overlap / narrow >= 0.25 or center_distance < image_width * 0.12:
                nearby.append(other)
        nearby.sort(key=lambda other: abs(other["cy"] - line["cy"]))
        context = " · ".join(dict.fromkeys(x["text"] for x in nearby[:4]))
        description = PRICE_RE.sub(" ", DATE_RE.sub(" ", raw))
        description = re.sub(r"\b(?:zł|zl|pln|z1)\b", " ", description, flags=re.I)
        text_chars = len(re.sub(r"[^a-zA-ZąćęłńóśźżĄĆĘŁŃÓŚŹŻ]", "", description))
        candidates.append({
            "page": page_no, "price": prices[0], "other_prices": prices[1:],
            "line": raw, "context": context, "ocr_confidence": line["confidence"],
            "x0": line["x0"], "x1": line["x1"],
            "y0": line["y0"], "y1": line["y1"], "cx": line["cx"], "cy": line["cy"], "description_chars": text_chars,
            "ambiguous_price_format": spaced_price,
        })
    return candidates


def extract_rows(pdf_path: Path) -> tuple[list[dict], int]:
    import fitz
    from PIL import Image, ImageEnhance, ImageFilter, ImageOps

    doc = fitz.open(pdf_path)
    rows: list[dict] = []
    for page_index, page in enumerate(doc, start=1):
        # 300 dpi retains the small price labels much better than the former 144 dpi pass.
        pix = page.get_pixmap(matrix=fitz.Matrix(300 / 72, 300 / 72), alpha=False)
        image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        gray = ImageOps.autocontrast(image.convert("L"))
        enhanced = ImageEnhance.Contrast(gray).enhance(1.35).filter(ImageFilter.UnsharpMask(radius=1.4, percent=145, threshold=3))
        binary = enhanced.point(lambda value: 255 if value > 178 else 0)
        pass_a = line_candidates(ocr_lines(enhanced, "pol+eng", "--psm 11"), page_index, image.width)
        pass_b = line_candidates(ocr_lines(binary, "pol+eng", "--psm 6"), page_index, image.width)

        merged = []
        used_b: set[int] = set()
        for a in pass_a:
            possible = [(abs(b["cy"] - (a["y0"] + a["y1"]) / 2), j, b) for j, b in enumerate(pass_b)
                        if j not in used_b and b["page"] == a["page"] and b["price"] == a["price"]
                        and abs(b["cx"] - (a["x0"] + a["x1"]) / 2) < image.width * 0.14
                        and abs(b["cy"] - (a["y0"] + a["y1"]) / 2) < 200]
            if possible:
                _, j, b = min(possible, key=lambda item: item[0])
                used_b.add(j)
                merged.append((a, b))
            else:
                merged.append((a, None))
        merged.extend((b, None) for j, b in enumerate(pass_b) if j not in used_b)

        for primary, corroborating in merged:
            line_scores = [primary["ocr_confidence"]]
            if corroborating:
                line_scores.append(corroborating["ocr_confidence"])
            line_confidence = round(sum(line_scores) / len(line_scores))
            description_chars = max(primary["description_chars"], corroborating["description_chars"] if corroborating else 0)
            corroborated = corroborating is not None
            ambiguous_price = primary["ambiguous_price_format"] or (corroborating["ambiguous_price_format"] if corroborating else False)
            if corroborated and line_confidence >= 80 and description_chars >= 10 and not ambiguous_price:
                level = "Powtórzony, mocny odczyt — sprawdź"
            elif corroborated and line_confidence >= 60 and not ambiguous_price:
                level = "Cena powtórzona — sprawdź"
            else:
                level = "Niepotwierdzony / niejednoznaczny — sprawdź"
            raw_line = primary["line"]
            if corroborating and corroborating["line"] != raw_line:
                raw_line += " | OCR 2: " + corroborating["line"]
            context = primary["context"] or (corroborating["context"] if corroborating else "")
            rows.append({
                "Sklep": "Biedronka",
                "Produkt / tekst rozpoznany": raw_line[:300],
                "Tekst w pobliżu — pomocniczy, może dotyczyć sąsiedniej oferty": context[:500],
                "Cena kandydująca (zł)": primary["price"],
                "Inne ceny z tej samej linii": ", ".join(f"{x:.2f}" for x in primary["other_prices"]),
                "Pewność OCR linii (%)": line_confidence,
                "Zgodność dwóch odczytów": "Tak" if corroborated else "Nie",
                "Ocena OCR": level,
                "Format ceny niejednoznaczny": "Tak — rozdzielenie cyfr przez OCR" if ambiguous_price else "Nie",
                "Strona": page_index,
                "Status": "Do porównania z gazetką — OCR nie potwierdza poprawności ceny",
                "Źródło": "",
            })
    return rows, len(doc)


def build_workbook(rows: list[dict], pdf_path: Path, pdf_url: str, leaflet_url: str, pages: int) -> Path:
    digest = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
    date_tag = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    output = OUT / f"biedronka-oferty-{date_tag}.xlsx"
    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    headers = list(rows[0]) if rows else ["Sklep", "Produkt / tekst rozpoznany", "Tekst w pobliżu — pomocniczy, może dotyczyć sąsiedniej oferty", "Cena kandydująca (zł)", "Inne ceny z tej samej linii", "Pewność OCR linii (%)", "Zgodność dwóch odczytów", "Ocena OCR", "Format ceny niejednoznaczny", "Strona", "Status", "Źródło"]
    widths = [18, 58, 55, 24, 25, 21, 23, 38, 32, 12, 54, 52]
    for row in rows:
        row["Źródło"] = leaflet_url

    def write_review_sheet(ws, subset: list[dict]):
        ws.append(headers)
        for item in subset:
            ws.append([item.get(header, "") for header in headers])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176B55")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 36
        for record in ws.iter_rows(min_row=2):
            for cell in record:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            record[3].number_format = '0.00 "zł"'

    higher = [row for row in rows if row["Ocena OCR"].startswith(("Powtórzony, mocny", "Cena powtórzona"))]
    low = [row for row in rows if row not in higher]
    write_review_sheet(wb.active, higher)
    wb.active.title = "Powtórzony OCR — sprawdź"
    write_review_sheet(wb.create_sheet("Niejednoznaczne — sprawdź"), low)
    info = wb.create_sheet("Informacje")
    info.append(["Gazetki → Excel — wynik automatyczny"])
    info.append(["Sklep", "Biedronka"])
    info.append(["Pobrano (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    info.append(["Stron PDF", pages])
    info.append(["Wierszy OCR", len(rows)])
    info.append(["Powtórzone, czytelniejsze odczyty", len(higher)])
    info.append(["Niepotwierdzone / niejednoznaczne", len(low)])
    info.append(["SHA-256 PDF", digest])
    info.append(["Strona gazetki", leaflet_url])
    info.append(["PDF źródłowy", pdf_url])
    info.append(["Ocena pewności", "To wskaźniki jakości OCR, nie prawdopodobieństwo poprawności oferty. Zgodność oznacza, że dwa ustawienia Tesseract na tym samym obrazie znalazły tę samą cenę w podobnym miejscu; kontekst może pochodzić z sąsiedniej oferty. Sprawdź wszystkie pozycje z PDF-em."])
    info.column_dimensions["A"].width = 23
    info.column_dimensions["B"].width = 110
    for cell in info[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=14)
        cell.fill = PatternFill("solid", fgColor="176B55")
    for row in info.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(output)
    (OUT / "source.json").write_text(json.dumps({"leaflet_url": leaflet_url, "pdf_url": pdf_url, "sha256": digest, "rows": len(rows)}, ensure_ascii=False, indent=2), encoding="utf-8")
    state_path = ROOT / ".state" / "biedronka.json"
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_path.write_text(json.dumps({"leaflet_url": leaflet_url, "pdf_url": pdf_url, "sha256": digest, "checked_at": datetime.now(timezone.utc).isoformat(timespec="seconds")}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return output


def main() -> int:
    detail_url, title, index_url = find_leaflet()
    log(f"Strona ofert: {index_url}\nWybrana gazetka: {detail_url}\nKarta: {title[:160]}")
    pdf_url = find_pdf(detail_url)
    pdf_path = download_pdf(pdf_url)
    log(f"Pobrano PDF: {pdf_path.stat().st_size} bajtów")
    digest = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
    state_path = ROOT / ".state" / "biedronka.json"
    if state_path.exists() and not os.getenv("FORCE_PROCESS"):
        try:
            previous = json.loads(state_path.read_text(encoding="utf-8"))
        except (ValueError, OSError):
            previous = {}
        if previous.get("sha256") == digest:
            log("Gazetka bez zmian — pomijam ponowny OCR.")
            if os.getenv("GITHUB_OUTPUT"):
                with open(os.environ["GITHUB_OUTPUT"], "a", encoding="utf-8") as output:
                    output.write("changed=false\n")
            return 0
    rows, pages = extract_rows(pdf_path)
    workbook = build_workbook(rows, pdf_path, pdf_url, detail_url, pages)
    log(f"Utworzono {workbook.name}: {len(rows)} wierszy OCR, {pages} stron")
    if os.getenv("GITHUB_OUTPUT"):
        with open(os.environ["GITHUB_OUTPUT"], "a", encoding="utf-8") as output:
            output.write(f"workbook={workbook.name}\n")
            output.write("changed=true\n")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        log(f"BŁĄD: {exc}")
        raise
